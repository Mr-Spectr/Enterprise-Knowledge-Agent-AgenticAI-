"""Role-filtered local RAG store for approved institutional documents.

The first implementation is intentionally dependency-light: it extracts text,
chunks it, persists metadata in SQLite, and uses lexical retrieval.  Its public
contract is stable so a vector/embedding backend can replace the ranking layer
without weakening access controls or source citations.
"""

from __future__ import annotations

import re
import sqlite3
from collections import Counter
from contextlib import contextmanager
from hashlib import sha256
from pathlib import Path
from typing import Iterable


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_DB_PATH = BASE_DIR / "data" / "knowledge.db"
_TOKEN = re.compile(r"[a-z0-9]{2,}")
_VISIBLE_LEVELS = {
    "unknown": {"public"},
    "student": {"public", "student"},
    "faculty": {"public", "student", "faculty"},
    "admin": {"public", "student", "faculty", "admin"},
}


@contextmanager
def _connect(db_path: Path = DEFAULT_DB_PATH) -> Iterable[sqlite3.Connection]:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    connection.execute(
        """CREATE TABLE IF NOT EXISTS knowledge_chunks (
            id TEXT PRIMARY KEY, document_name TEXT NOT NULL, source_path TEXT NOT NULL,
            page_number INTEGER, chunk_index INTEGER NOT NULL, access_level TEXT NOT NULL,
            content TEXT NOT NULL, content_hash TEXT NOT NULL)"""
    )
    connection.execute("CREATE INDEX IF NOT EXISTS idx_knowledge_access ON knowledge_chunks(access_level)")
    connection.commit()
    try:
        yield connection
    finally:
        connection.close()


def _tokens(text: str) -> list[str]:
    return _TOKEN.findall(text.lower())


def _chunks(text: str, size: int = 1000, overlap: int = 150) -> Iterable[str]:
    clean = re.sub(r"\s+", " ", text).strip()
    start = 0
    while start < len(clean):
        end = min(len(clean), start + size)
        if end < len(clean):
            split = clean.rfind(" ", start, end)
            end = split if split > start + size // 2 else end
        value = clean[start:end].strip()
        if value:
            yield value
        if end >= len(clean):
            return
        start = max(start + 1, end - overlap)


def _read_document(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return [(None, path.read_text(encoding="utf-8", errors="replace"))]
    if suffix == ".pdf":
        from pypdf import PdfReader
        return [(index + 1, page.extract_text() or "") for index, page in enumerate(PdfReader(path).pages)]
    if suffix == ".docx":
        from docx import Document
        return [(None, "\n".join(item.text for item in Document(path).paragraphs))]
    raise ValueError("Only PDF, DOCX, TXT, and Markdown documents can be indexed.")


def ingest_document(path: str | Path, access_level: str = "public", db_path: Path = DEFAULT_DB_PATH) -> dict:
    source = Path(path).resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Document not found: {source}")
    if access_level not in {"public", "student", "faculty", "admin"}:
        raise ValueError("Invalid document access level.")
    rows = []
    for page, text in _read_document(source):
        for index, content in enumerate(_chunks(text)):
            identity = sha256(f"{source}|{page}|{index}|{content}".encode()).hexdigest()
            rows.append((identity, source.name, str(source), page, index, access_level, content, sha256(content.encode()).hexdigest()))
    with _connect(db_path) as connection:
        connection.execute("DELETE FROM knowledge_chunks WHERE source_path = ?", (str(source),))
        connection.executemany(
            "INSERT INTO knowledge_chunks VALUES (?, ?, ?, ?, ?, ?, ?, ?)", rows
        )
        connection.commit()
    return {"document": source.name, "chunks_indexed": len(rows), "access_level": access_level}


def search_knowledge(query: str, requester_role: str, limit: int = 5, db_path: Path = DEFAULT_DB_PATH) -> dict:
    query_terms = Counter(_tokens(query))
    if not query_terms:
        raise ValueError("Query must contain searchable words.")
    levels = sorted(_VISIBLE_LEVELS.get(requester_role, {"public"}))
    placeholders = ", ".join("?" for _ in levels)
    with _connect(db_path) as connection:
        candidates = connection.execute(
            f"SELECT document_name, page_number, chunk_index, content FROM knowledge_chunks WHERE access_level IN ({placeholders})",
            levels,
        ).fetchall()
    ranked = []
    for row in candidates:
        counts = Counter(_tokens(row["content"]))
        score = sum(counts[term] * weight for term, weight in query_terms.items())
        if score:
            ranked.append((score, row))
    ranked.sort(key=lambda item: (-item[0], item[1]["document_name"], item[1]["chunk_index"]))
    results = [
        {"source": row["document_name"], "page": row["page_number"], "chunk": row["chunk_index"],
         "excerpt": row["content"][:700], "score": score}
        for score, row in ranked[:max(1, min(limit, 10))]
    ]
    return {"query": query, "role": requester_role, "result_count": len(results), "results": results}


def bootstrap_packaged_knowledge() -> list[dict]:
    """Idempotently index docs distributed with the application image."""
    corpus = BASE_DIR / "knowledge"
    if not corpus.is_dir():
        return []
    return [ingest_document(path, "public") for path in sorted(corpus.glob("*")) if path.suffix.lower() in {".md", ".txt", ".pdf", ".docx"}]
